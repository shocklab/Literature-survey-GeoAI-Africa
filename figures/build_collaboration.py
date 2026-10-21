#!/usr/bin/env python3
"""
Generate fig-collaboration.tex (TikZ/pgfplots) from the reviewed-studies workbook,
house style (style.tex colours, jsxbar + manual-arc pies). No extra LaTeX packages.

    python3 build_collaboration.py [source/Reviewed_articles_only.xlsx] [out.tex]

Layout: (a) contributor-country bars on the left; (b) leadership pie over
(c) non-academic partner-type pie on the right, each with a right-hand legend.
All percentages rendered to 1 decimal place.
"""
import sys, collections
from pathlib import Path
import openpyxl

HERE = Path(__file__).resolve().parent
SRC  = Path(sys.argv[1]) if len(sys.argv) > 1 else HERE / "source" / "Reviewed_articles_only.xlsx"
OUT  = Path(sys.argv[2]) if len(sys.argv) > 2 else HERE / "fig-collaboration.tex"
TOPN = 13

AFRICA = {'DZ','AO','BJ','BW','BF','BI','CM','CV','CF','TD','KM','CG','CD','CI','DJ','EG',
 'GQ','ER','SZ','ET','GA','GM','GH','GN','GW','KE','LS','LR','LY','MG','MW','ML','MR','MU',
 'MA','MZ','NA','NE','NG','RW','ST','SN','SC','SL','SO','ZA','SS','SD','TZ','TG','TN','UG','ZM','ZW'}
ISO2NAME = {'US':'United States','ZA':'South Africa','DE':'Germany','GB':'United Kingdom',
 'IT':'Italy','CN':'China','NL':'Netherlands','KE':'Kenya','BE':'Belgium','FR':'France',
 'NG':'Nigeria','CA':'Canada','GH':'Ghana','ES':'Spain','ET':'Ethiopia','IN':'India'}

wb = openpyxl.load_workbook(SRC, data_only=True); ws = wb["Studies"]
hdr = [c.value for c in ws[1]]; idx = {h: i for i, h in enumerate(hdr)}
def cell(name, r): return ws.cell(r, idx[name]+1).value

N = 0; ctry = collections.Counter(); lead = collections.Counter()
for r in range(2, ws.max_row+1):
    if cell("study_id", r) is None: continue
    N += 1
    ic = str(cell("inst_countries", r) or "")
    if ic and ic != "Unknown":
        for c in [x.strip().upper() for x in ic.replace(",", ";").split(";") if x.strip()]:
            ctry[c] += 1
    l = cell("african_leadership", r)
    if l and l != "Unknown": lead[l] += 1

top  = ctry.most_common(TOPN)[::-1]
def nm(c): return ISO2NAME.get(c, c)
ypos = {c: i+1 for i, (c, _) in enumerate(top)}
n_a  = len(top)
ylabs  = ",".join("{%s}" % nm(c) for c, _ in top)
yticks = ",".join(str(i+1) for i in range(n_a))
afr_coords = " ".join("(%d,%d)" % (v, ypos[c]) for c, v in top if c in AFRICA)
non_coords = " ".join("(%d,%d)" % (v, ypos[c]) for c, v in top if c not in AFRICA)
xmax   = (max(v for _, v in top)//10 + 1)*10 + 8
xticks = ",".join(str(t) for t in range(0, xmax-6, 10))
a_h    = round(0.46*len(top) + 0.5, 1)

TPL_PIE = r"""\begin{tikzpicture}
  \def\R{1.25}
  \tikzset{
    slice/.style={draw=white, line width=1.1pt},
    inpct/.style={font=\scriptsize},
    legtxt/.style={font=\scriptsize, text=jsAxis},
  }
  \node[font=\small\bfseries, text=jsAxis, anchor=south west] at (-1.5,1.45) {@LBL@};
@BODY@
\end{tikzpicture}"""

def make_pie(label, entries):
    tot = sum(c for _, c, _ in entries)
    k = len(entries); step = 0.52; y0 = (k-1)/2.0*step
    wedges, inpct, legend = [], [], []
    ang = 90.0
    for i, (name, c, col) in enumerate(entries):
        sweep = c/tot*360.0
        a0, a1 = ang, ang - sweep
        mid = (a0 + a1)/2.0
        wedges.append(r"  \filldraw[slice, fill=%s] (0,0) -- (%.2f:\R) arc (%.2f:%.2f:\R) -- cycle;"
                      % (col, a0, a0, a1))
        pct = c/tot*100
        dark = col in ("jsGrey", "jsAmber")
        tcol = "jsAxis" if dark else "white"
        if sweep >= 28:
            inpct.append(r"  \node[inpct, text=%s] at (%.2f:0.60*\R) {%.1f\%%};" % (tcol, mid, pct))
            legtxt = r"%s ($n=%d$)" % (name, c)
        else:
            legtxt = r"%s (%.1f\%%, $n=%d$)" % (name, pct, c)
        yy = y0 - i*step
        legend.append(r"  \filldraw[fill=%s, draw=white, line width=0.4pt] (2.05,%.2f) rectangle ++(0.28,0.28);"
                      % (col, yy-0.14))
        legend.append(r"  \node[legtxt, anchor=west] at (2.45,%.2f) {%s};" % (yy, legtxt))
        ang = a1
    body = "\n".join(wedges) + "\n" + "\n".join(inpct) + "\n" + "\n".join(legend)
    return TPL_PIE.replace("@LBL@", label).replace("@BODY@", body)

B_ENTRIES = [("African-only",     lead.get("African-only", 0),     "jsGreen"),
             ("North--South",     lead.get("North-South", 0),      "jsBlue"),
             ("Non-African-only", lead.get("Non-African-only", 0), "jsGrey")]
# (c) non-academic partner-type counts among the 70 cross-sector studies (OpenAlex
# 'type' of each non-academic partner; partner instances, a study may contribute
# several). The 4 manually-coded studies are education-only. Regenerate via the
# affiliation script if the corpus changes.
C_ENTRIES = [("Research facility", 86, "jsTeal"),
             ("Non-profit / NGO",  33, "jsAmber"),
             ("Government",        32, "jsPlum"),
             ("Company",           10, "jsRust")]
N_CROSS = 70

pie_b = make_pie("(b)", B_ENTRIES)
pie_c = make_pie("(c)", C_ENTRIES)
afr_pct = 100*(lead.get("African-only", 0)+lead.get("North-South", 0))/N
non_pct = 100*lead.get("Non-African-only", 0)/N

BANNER = ("%% GENERATED by build_collaboration.py from source/Reviewed_articles_only.xlsx\n"
          "%% Do not edit by hand: edit the spreadsheet and re-run the script.\n")

TPL = r"""
\begin{figure*}[pos=!t]
\centering
\begin{minipage}[c]{0.50\textwidth}
\centering
\begin{tikzpicture}
\begin{axis}[
  jsxbar, scale only axis, width=4.8cm, height=@AH@cm,
  title={(a)}, area legend,
  xlabel={Studies with an author institution in the country},
  xmin=0, xmax=@XMAX@, xtick={@XTICKS@},
  ymin=0.4, ymax=@NA@.6, ytick={@YTICKS@}, yticklabels={@YLABS@},
  bar width=7pt, enlarge y limits=0,
  nodes near coords={\pgfmathprintnumber[fixed,precision=0]{\pgfplotspointmeta}},
  legend style={at={(0.98,0.02)}, anchor=south east, fill=white, fill opacity=0.92, text opacity=1, font=\footnotesize},
  legend cell align=left,
]
\addplot[draw=none, fill=jsGreen, bar shift=0pt] coordinates {@AFR@};
\addplot[draw=none, fill=jsGrey,  bar shift=0pt] coordinates {@NON@};
\legend{African, Non-African}
\end{axis}
\end{tikzpicture}
\end{minipage}\hfill
\begin{minipage}[c]{0.48\textwidth}
\centering
@PIEB@

\vspace{1.0em}

@PIEC@
\end{minipage}
\caption{Institutional collaboration across the reviewed studies, from author affiliations. \textbf{(a)}~Contributing institution countries, the number of studies with at least one author institution in each country, top 13 shown. \textbf{(b)}~Share of studies by African institutional leadership, with collaboration between African-only, North--South, and non-African-only institutions. \textbf{(c)}~Non-academic partners of the studies that combine academic and non-academic institutions, by institution type.}
\label{fig:collaboration}
\end{figure*}
"""
rep = {"@AH@": "%.1f" % a_h, "@XMAX@": str(xmax), "@XTICKS@": xticks, "@NA@": str(n_a),
       "@YTICKS@": yticks, "@YLABS@": ylabs, "@AFR@": afr_coords, "@NON@": non_coords,
       "@N@": str(N), "@TOPN@": str(len(top)), "@PIEB@": pie_b, "@PIEC@": pie_c,
       "@AFRP@": "%.1f" % afr_pct, "@NONP@": "%.1f" % non_pct, "@NCROSS@": str(N_CROSS)}
tex = TPL
for k, v in rep.items(): tex = tex.replace(k, v)
OUT.write_text(BANNER + tex)
print("wrote %s | N=%d | African-inst %.1f%% / none %.1f%%" % (OUT, N, afr_pct, non_pct))
